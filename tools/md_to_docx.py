"""Convert docs/API_REFERENCE.md to PatchSignal-API-Reference.docx. Run: python3 tools/md_to_docx.py"""
import re
from docx import Document
from docx.shared import Pt, RGBColor

SRC = "docs/API_REFERENCE.md"
OUT = "PatchSignal-API-Reference.docx"
CODE_FONT = "Consolas"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)

def add_runs(para, text):
    parts = re.split(r"(\*\*.+?\*\*|`.+?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = para.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = para.add_run(part[1:-1]); r.font.name = CODE_FONT; r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x5F)
        else:
            para.add_run(part)

lines = open(SRC, encoding="utf-8").read().splitlines()
i = 0
in_code = False
code_buf: list[str] = []
table_buf: list[list[str]] = []

def flush_code():
    global code_buf
    if code_buf:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run("\n".join(code_buf))
        r.font.name = CODE_FONT; r.font.size = Pt(9)
        code_buf = []

def flush_table():
    global table_buf
    rows = [r for r in table_buf if not all(re.fullmatch(r":?-{3,}:?", c.strip()) for c in r)]
    if rows:
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Light Grid Accent 1"
        for ri, row in enumerate(rows):
            for ci, cell in enumerate(row):
                t.cell(ri, ci).text = cell.strip().strip("`")
                for p in t.cell(ri, ci).paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9.5)
    table_buf = []

while i < len(lines):
    line = lines[i]
    if line.strip().startswith("```"):
        if in_code:
            flush_code(); in_code = False
        else:
            flush_table(); in_code = True
        i += 1
        continue
    if in_code:
        code_buf.append(line); i += 1
        continue
    if line.strip().startswith("|") and line.strip().endswith("|"):
        table_buf.append([c for c in line.strip().strip("|").split("|")])
        i += 1
        continue
    else:
        flush_table()
    s = line.strip()
    if not s:
        i += 1
        continue
    if s.startswith("# "):
        doc.add_heading(s[2:], level=0)
    elif s.startswith("## "):
        doc.add_heading(s[3:], level=1)
    elif s.startswith("### "):
        doc.add_heading(s[4:], level=2)
    elif s.startswith("> "):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(18)
        r = p.add_run(s[2:]); r.italic = True
    elif re.match(r"^(\d+)\.\s", s):
        p = doc.add_paragraph(style="List Number")
        add_runs(p, re.sub(r"^\d+\.\s", "", s))
    elif s.startswith("- ") or s.startswith("* "):
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, s[2:])
    else:
        p = doc.add_paragraph()
        add_runs(p, s)
    i += 1

flush_code(); flush_table()
doc.save(OUT)
print(f"Saved {OUT}")
